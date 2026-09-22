"""docbridge · docx 管线离线单元测试（禁止联网）。

用法：
    cd <仓库目录>/docbridge
    python3 -m pytest tests/test_docx_pipeline.py -q      # 若有 pytest
    python3 tests/test_docx_pipeline.py                   # 无 pytest 也能跑

样例文档全部由 python-docx 现场生成，写入 tests/_tmp/。
"""

from __future__ import annotations

import importlib.util
import base64
import os
import sys
import traceback

from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
TMP = os.path.join(HERE, "_tmp")
os.makedirs(TMP, exist_ok=True)


def _load_pipeline():
    """按文件路径加载 pipelines/docx_pipeline.py，避免依赖他人维护的 __init__.py。"""
    path = os.path.join(ROOT, "pipelines", "docx_pipeline.py")
    spec = importlib.util.spec_from_file_location("docx_pipeline_under_test", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


docx_pipeline = _load_pipeline()

TRANSLATED_PREFIX = "【译"
PURE_PREFIX = "【译文"

# 1x1 像素 PNG（离线生成图片用，避免依赖外部素材或联网）
_PIXEL_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8DwHwAFAAH/"
    "q842iQAAAABJRU5ErkJggg=="
)


def pixel_png_path():
    path = os.path.join(TMP, "pixel.png")
    if not os.path.exists(path):
        with open(path, "wb") as fh:
            fh.write(base64.b64decode(_PIXEL_PNG_B64))
    return path


# ---------------------------------------------------------------- 假翻译函数


def fake(texts, ctx=None):
    """契约里的标准假翻译：保留原文并加前缀（便于检查插入位置）。"""
    assert isinstance(texts, list)
    assert ctx is None or isinstance(ctx, dict)
    return [f"【译{i}】{t}" for i, t in enumerate(texts)]


def fake_pure(texts, ctx=None):
    """彻底替换原文（便于断言“原文特征串消失”）。"""
    return [f"【译文{i}】" for i in range(len(texts))]


def fake_short(texts, ctx=None):
    """故意少返回一条，run 必须抛 RuntimeError。"""
    return fake(texts)[:-1]


class _Recorder:
    """记录 translate 调用批次与 progress 回调。"""

    def __init__(self, fn=fake):
        self.fn = fn
        self.batches = []
        self.progress_calls = []

    def translate(self, texts, ctx=None):
        self.batches.append(list(texts))
        return self.fn(texts, ctx)

    def progress(self, done, total, note=""):
        self.progress_calls.append((done, total, note))


def assert_raises(exc_type, fn, message_part=None):
    try:
        fn()
    except exc_type as exc:  # noqa: PERF203
        if message_part is not None and message_part not in str(exc):
            raise AssertionError(f"异常信息里没有 {message_part!r}：{exc}")
        return exc
    except Exception as exc:  # noqa: BLE001
        raise AssertionError(f"期望 {exc_type.__name__}，实际 {type(exc).__name__}: {exc}")
    raise AssertionError(f"期望抛出 {exc_type.__name__}，但没有任何异常")


# ---------------------------------------------------------------- 样例文档


def make_sample(path, body_paragraphs=None):
    """生成含标题/正文/表格(含嵌套)/页眉/页脚的样例文档。"""
    doc = Document()
    doc.add_heading("Introduction", level=1)
    if body_paragraphs is None:
        body_paragraphs = [
            "This is the first body paragraph.",
            "12345",
            "这是中文段落。",
            "Second body paragraph with more words.",
        ]
    for text in body_paragraphs:
        doc.add_paragraph(text)

    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Table cell one"
    table.cell(0, 1).text = "Table cell two"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested cell text"

    header = doc.sections[0].header
    header.is_linked_to_previous = False
    header.paragraphs[0].text = "Header Title"

    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    footer.paragraphs[0].text = "Footer Text"

    doc.save(path)
    return path


def part_texts(doc):
    """按分区返回文本：body/table/header/footer 的全部段落文本（文档顺序）。"""
    out = {"body": [], "table": [], "header": [], "footer": []}

    def walk(container, key, depth=0):
        for p in container.paragraphs:
            out[key].append(p.text)
        if depth >= 6:
            return
        for t in container.tables:
            for row in t.rows:
                for cell in row.cells:
                    walk(cell, "table", depth + 1)

    walk(doc, "body")
    for section in doc.sections:
        for attr, key in (("header", "header"), ("footer", "footer")):
            container = getattr(section, attr)
            if container.is_linked_to_previous:
                continue
            for p in container.paragraphs:
                out[key].append(p.text)
    return out


def flat_text(doc):
    parts = part_texts(doc)
    return "\n".join(parts["body"] + parts["table"] + parts["header"] + parts["footer"])


def paragraph_count(doc):
    parts = part_texts(doc)
    return sum(len(v) for v in parts.values())


def style_names(doc):
    names = set()

    def walk(container, depth=0):
        for p in container.paragraphs:
            names.add(p.style.name)
        if depth >= 6:
            return
        for t in container.tables:
            for row in t.rows:
                for cell in row.cells:
                    walk(cell, depth + 1)

    walk(doc, 0)
    for section in doc.sections:
        for attr in ("header", "footer"):
            container = getattr(section, attr)
            if container.is_linked_to_previous:
                continue
            for p in container.paragraphs:
                names.add(p.style.name)
    return names


def check_progress(calls, expected_total):
    assert calls, "progress 回调一次都没被调用"
    dones = [c[0] for c in calls]
    totals = [c[1] for c in calls]
    assert all(b >= a for a, b in zip(dones, dones[1:])), f"progress 不是单调递增：{dones}"
    assert all(t == totals[0] for t in totals), f"progress 的 total 不稳定：{totals}"
    assert totals[-1] == expected_total, f"total 应为 {expected_total}，实际 {totals[-1]}"
    assert dones[-1] == totals[-1], f"最后一次 done({dones[-1]}) != total({totals[-1]})"


# ---------------------------------------------------------------- 契约与导出


def test_module_exports_contract():
    assert docx_pipeline.FORMAT == "docx"
    assert docx_pipeline.MODE == "both"
    assert isinstance(docx_pipeline.LABEL, str) and docx_pipeline.LABEL
    assert isinstance(docx_pipeline.NOTE, str) and docx_pipeline.NOTE
    assert isinstance(docx_pipeline.OPTIONS, list) and docx_pipeline.OPTIONS
    assert callable(docx_pipeline.run)


def test_run_signature_is_keyword_only():
    import inspect

    sig = inspect.signature(docx_pipeline.run)
    params = sig.parameters
    assert list(params)[:2] == ["src_path", "out_path"]
    for name in ("translate", "progress", "options"):
        assert params[name].kind is inspect.Parameter.KEYWORD_ONLY, f"{name} 必须是关键字参数"


# ---------------------------------------------------------------- 原位替换


def test_inplace_replaces_text_and_keeps_structure():
    src = make_sample(os.path.join(TMP, "inplace_src.docx"))
    out = os.path.join(TMP, "inplace_out.docx")
    if os.path.exists(out):
        os.remove(out)

    rec = _Recorder(fake_pure)
    result = docx_pipeline.run(
        src, out, translate=rec.translate, progress=rec.progress,
        options={"__mode__": "inplace", "target_lang": "zh-Hans"},
    )

    assert os.path.exists(out), "输出文件不存在"
    doc = Document(out)  # 必须能重新打开
    text = flat_text(doc)

    # 原文特征串消失、译文特征串出现
    for original in ("Introduction", "This is the first body paragraph.",
                     "Second body paragraph with more words.",
                     "Table cell one", "Table cell two", "Nested cell text",
                     "Header Title", "Footer Text"):
        assert original not in text, f"原位模式下原文仍在：{original}"
    assert f"{PURE_PREFIX}0】" in text and PURE_PREFIX in text

    # 段落数与样式数不变
    assert paragraph_count(doc) == paragraph_count(Document(src))
    assert style_names(doc) == style_names(Document(src))

    # 页眉也翻译了
    assert TRANSLATED_PREFIX in part_texts(doc)["header"][0] or PURE_PREFIX in part_texts(doc)["header"][0]
    assert PURE_PREFIX in part_texts(doc)["header"][0], "页眉没有被翻译"

    # 跳过统计：空/数字/中文
    assert result["units"] >= 6, result
    assert result["skipped"] >= 2, result  # 12345 与中文段落
    assert result["chars_in"] > 0 and result["chars_out"] > 0
    assert result["mode"] == "inplace"
    assert isinstance(result["detail"], str) and result["detail"]

    check_progress(rec.progress_calls, result["units"])


def test_inplace_skips_digits_and_chinese_untouched():
    src = make_sample(os.path.join(TMP, "inplace_skip_src.docx"))
    out = os.path.join(TMP, "inplace_skip_out.docx")

    rec = _Recorder(fake)
    result = docx_pipeline.run(src, out, translate=rec.translate, progress=rec.progress,
                               options={"__mode__": "inplace"})

    doc = Document(out)
    body = part_texts(doc)["body"]
    assert "12345" in body, "纯数字段落被改动了"
    assert "这是中文段落。" in body, "中文段落被改动了"
    assert result["skipped"] >= 2
    # 数字与中文段落没有被送进翻译
    sent = [t for batch in rec.batches for t in batch]
    assert "12345" not in sent
    assert "这是中文段落。" not in sent

    # 译文 run 设置了东亚字体（eastAsia），且没有强改 Latin 字体（ascii）
    translated_runs = []
    for p in doc.paragraphs:
        for r in p.runs:
            if TRANSLATED_PREFIX in r.text:
                translated_runs.append(r)
    assert translated_runs, "找不到译文 run"
    for r in translated_runs:
        rfonts = r._r.rPr.rFonts if r._r.rPr is not None else None
        assert rfonts is not None and rfonts.get(qn("w:eastAsia")), "译文 run 缺少 w:eastAsia 字体"

    check_progress(rec.progress_calls, result["units"])


# ---------------------------------------------------------------- 双语对照


def test_bilingual_inserts_translation_right_after_original():
    src = make_sample(os.path.join(TMP, "bilingual_src.docx"))
    out = os.path.join(TMP, "bilingual_out.docx")

    rec = _Recorder(fake)
    result = docx_pipeline.run(src, out, translate=rec.translate, progress=rec.progress,
                               options={"__mode__": "bilingual"})

    assert os.path.exists(out)
    doc = Document(out)
    paras = doc.paragraphs
    texts = [p.text for p in paras]

    # 原文仍在
    assert "This is the first body paragraph." in texts
    assert "Introduction" in texts
    assert "这是中文段落。" in texts and "12345" in texts

    # 译文段落紧邻出现在原文段落之后
    checked = 0
    for i, t in enumerate(texts):
        if t.startswith(TRANSLATED_PREFIX):
            continue
        if not t.strip() or t.strip() == "12345" or t.strip() == "这是中文段落。":
            continue
        nxt = texts[i + 1] if i + 1 < len(texts) else ""
        assert nxt.startswith(TRANSLATED_PREFIX), f"段落 {t!r} 之后没有紧邻译文，而是 {nxt!r}"
        assert t in nxt, "译文段应包含对应的原文（fake 翻译加前缀）"
        checked += 1
    assert checked >= 3, f"检查到的双语对太少：{checked}"

    # 译文段落被标成深蓝色（非高亮）
    color_found = False
    for p in paras:
        if p.text.startswith(TRANSLATED_PREFIX):
            for r in p.runs:
                rgb = r.font.color.rgb if r.font.color and r.font.color.type is not None else None
                if rgb == RGBColor(0x1F, 0x4E, 0x79):
                    color_found = True
                assert r.font.highlight_color is None, "译文使用了高亮"
    assert color_found, "译文段落没有设置深蓝色"

    # 表格单元格与页眉也插入了译文
    table_texts = part_texts(doc)["table"]
    assert any(t.startswith(TRANSLATED_PREFIX) for t in table_texts), "表格单元格没有插入译文"
    assert any("Nested cell text" in t for t in table_texts), "嵌套表格原文丢失"
    header_texts = part_texts(doc)["header"]
    assert len(header_texts) == 2 and header_texts[1].startswith(TRANSLATED_PREFIX), header_texts
    footer_texts = part_texts(doc)["footer"]
    assert len(footer_texts) == 2 and footer_texts[1].startswith(TRANSLATED_PREFIX), footer_texts

    # 双语模式下段落数增加（每段原文后加一段）
    assert paragraph_count(doc) == paragraph_count(Document(src)) + result["units"]

    check_progress(rec.progress_calls, result["units"])


def test_bilingual_heading_style_is_inherited():
    src = make_sample(os.path.join(TMP, "bilingual_style_src.docx"))
    out = os.path.join(TMP, "bilingual_style_out.docx")

    docx_pipeline.run(src, out, translate=fake, options={"__mode__": "bilingual"})
    doc = Document(out)
    paras = doc.paragraphs
    assert paras[0].text == "Introduction" and paras[0].style.name == "Heading 1"
    assert paras[1].text.startswith(TRANSLATED_PREFIX)
    assert paras[1].style.name == "Heading 1", "译文段落没有继承原标题样式"
    # 译文段落不应复制出分节符
    for p in paras:
        assert p._p.find(qn("w:pPr")) is None or p._p.find(qn("w:pPr")).find(qn("w:sectPr")) is None


# ---------------------------------------------------------------- 批量化 / 进度


def test_translate_is_batched_between_30_and_100():
    paragraphs = [f"Paragraph number {i} of the batch test." for i in range(150)]
    src = make_sample(os.path.join(TMP, "batch_src.docx"), body_paragraphs=paragraphs)
    out = os.path.join(TMP, "batch_out.docx")

    rec = _Recorder(fake)
    result = docx_pipeline.run(src, out, translate=rec.translate, progress=rec.progress,
                               options={"__mode__": "inplace"})

    sizes = [len(b) for b in rec.batches]
    # 150 个正文段落 + 标题/表格(含嵌套)/页眉/页脚共 6 段
    assert result["units"] == 156, result
    assert len(rec.batches) >= 2, f"没有攒批调用，实际批次：{sizes}"
    assert max(sizes) <= 100, f"单批超过 100 条：{sizes}"
    assert sum(sizes) == result["units"]
    # 批大小配置生效且被夹到 30~100
    rec2 = _Recorder(fake)
    docx_pipeline.run(src, os.path.join(TMP, "batch_out2.docx"), translate=rec2.translate,
                      options={"__mode__": "inplace", "batch_size": 5})
    assert max(len(b) for b in rec2.batches) >= 30, "batch_size 没有夹到合理下限"

    check_progress(rec.progress_calls, result["units"])


def test_progress_final_done_equals_total_when_all_skipped():
    doc = Document()
    doc.add_paragraph("12345")
    doc.add_paragraph("")
    src = os.path.join(TMP, "empty_src.docx")
    doc.save(src)
    out = os.path.join(TMP, "empty_out.docx")

    rec = _Recorder(fake)
    result = docx_pipeline.run(src, out, translate=rec.translate, progress=rec.progress,
                               options={"__mode__": "inplace"})

    assert result["units"] == 0
    assert rec.batches == [], "没有可翻译单元却调用了 translate"
    assert os.path.exists(out)
    check_progress(rec.progress_calls, 0)


def test_paragraphs_with_picture_or_field_are_skipped():
    """含图片/域代码的段落不能动，否则会破坏它们。"""
    from docx.oxml import OxmlElement

    doc = Document()
    doc.add_paragraph("A normal sentence that should be translated.")

    picture_p = doc.add_paragraph()
    picture_p.add_run().add_picture(pixel_png_path())

    field_p = doc.add_paragraph()
    field_p.add_run("Total pages: ")
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    field_p.add_run()._r.append(fld)
    field_p.add_run("FIELDCODE")
    field_text_before = field_p.text

    src = os.path.join(TMP, "locked_src.docx")
    out = os.path.join(TMP, "locked_out.docx")
    doc.save(src)

    rec = _Recorder(fake)
    result = docx_pipeline.run(src, out, translate=rec.translate, progress=rec.progress,
                               options={"__mode__": "inplace"})
    assert result["skipped"] >= 2, result

    out_doc = Document(out)
    drawings_before = len(Document(src).element.body.findall(".//" + qn("w:drawing")))
    drawings_after = len(out_doc.element.body.findall(".//" + qn("w:drawing")))
    assert drawings_after == drawings_before == 1, "图片数量变了"
    assert out_doc.paragraphs[2].text == field_text_before, "含域代码的段落被改动了"
    assert out_doc.paragraphs[1].text == "", "含图片的段落被写入了文本"
    # 这些段落没有被送进翻译
    sent = [t for batch in rec.batches for t in batch]
    assert all("FIELDCODE" not in t for t in sent)
    assert result["units"] == 1, result


# ---------------------------------------------------------------- 失败路径


def test_translate_length_mismatch_raises_and_leaves_no_part():
    src = make_sample(os.path.join(TMP, "short_src.docx"))
    out = os.path.join(TMP, "short_out.docx")
    for path in (out, out + ".part"):
        if os.path.exists(path):
            os.remove(path)

    rec = _Recorder(fake_short)
    error = assert_raises(
        RuntimeError,
        lambda: docx_pipeline.run(src, out, translate=rec.translate, progress=rec.progress,
                                  options={"__mode__": "inplace"}),
        "翻译返回数量不符",
    )
    assert "期望" in str(error) and "实际" in str(error)
    assert not os.path.exists(out), "失败却写出了一个输出文件"
    assert not os.path.exists(out + ".part"), "失败留下了 .part 半成品"


def test_bad_inputs_raise_chinese_runtime_error():
    out = os.path.join(TMP, "never.docx")
    assert_raises(RuntimeError, lambda: docx_pipeline.run(os.path.join(TMP, "nope.docx"), out,
                                                          translate=fake), "源文件不存在")

    broken = os.path.join(TMP, "broken.docx")
    with open(broken, "wb") as fh:
        fh.write(b"this is not a docx file")
    assert_raises(RuntimeError, lambda: docx_pipeline.run(broken, out, translate=fake),
                  "无法打开 Word 文档")

    src = make_sample(os.path.join(TMP, "mode_src.docx"))
    assert_raises(RuntimeError, lambda: docx_pipeline.run(src, out, translate=fake,
                                                          options={"__mode__": "sideways"}),
                  "未知的 docx 处理模式")


def test_source_file_is_not_modified():
    src = make_sample(os.path.join(TMP, "readonly_src.docx"))
    before = os.path.getsize(src)
    with open(src, "rb") as fh:
        before_bytes = fh.read()
    docx_pipeline.run(src, os.path.join(TMP, "readonly_out.docx"), translate=fake,
                      options={"__mode__": "bilingual"})
    with open(src, "rb") as fh:
        after_bytes = fh.read()
    assert before == os.path.getsize(src) and before_bytes == after_bytes, "源文件被改动了"


def test_compatible_option_keys_do_not_crash():
    src = make_sample(os.path.join(TMP, "opts_src.docx"))
    out = os.path.join(TMP, "opts_out.docx")
    result = docx_pipeline.run(
        src, out, translate=fake,
        options={
            "__mode__": "inplace",
            "source_lang": "en",
            "target_lang": "zh-Hans",
            "font_scale": 0.9,
            "keep_original_style": False,
            "unknown_key_should_be_ignored": object(),
        },
    )
    assert os.path.exists(out) and result["units"] >= 6
    # mode / bilingual 也能识别
    r2 = docx_pipeline.run(src, os.path.join(TMP, "opts_out2.docx"), translate=fake,
                           options={"bilingual": True})
    assert r2["mode"] == "bilingual"


# ---------------------------------------------------------------- 独立运行


def _run_all():
    tests = [(name, obj) for name, obj in sorted(globals().items())
             if name.startswith("test_") and callable(obj)]
    failed = []
    for name, fn in tests:
        try:
            fn()
        except Exception:  # noqa: BLE001
            failed.append(name)
            print(f"FAIL {name}")
            traceback.print_exc()
        else:
            print(f"PASS {name}")
    print(f"\n共 {len(tests)} 个用例，失败 {len(failed)} 个")
    if failed:
        print("FAILED: " + ", ".join(failed))
        return 1
    print("OK")
    return 0


if __name__ == "__main__":
    sys.exit(_run_all())
